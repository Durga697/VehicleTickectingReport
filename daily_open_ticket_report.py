"""Download the current month's five ticket exports and build one open-ticket report."""

from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import sync_playwright

import download_vehicle_ticketing as downloader


ROOT = Path(__file__).resolve().parent
DOWNLOAD_DIR = ROOT / "Whizzard_Downloads"
REPORT_DIR = DOWNLOAD_DIR / "Daily_Open_Ticket_Reports"
LOG_DIR = DOWNLOAD_DIR / "_daily_logs"
REPORT_PATH = REPORT_DIR / f"Current_Month_Open_Tickets_{datetime.now():%d-%m-%y}.xlsx"

BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
RED = "C00000"
LIGHT_RED = "FCE4D6"
GREEN = "548235"
WHITE = "FFFFFF"
COMPLETED_STATUSES = {"accepted", "auto-accepted", "closed"}


def setup_logging() -> Path:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    log_path = LOG_DIR / f"daily_open_tickets_{datetime.now():%Y-%m-%d}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8"), logging.StreamHandler()],
    )
    return log_path


def download_current_month() -> list[dict[str, str | int]]:
    target = (datetime.now().year, datetime.now().month)
    downloader.OUTPUT_PATH = DOWNLOAD_DIR
    downloader.HEADLESS = True
    downloader.OVERWRITE_EXISTING = True
    debug_dir = downloader.prepare_folders()
    results: list[dict[str, str | int]] = []

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage"],
        )
        context = browser.new_context(
            accept_downloads=True,
            viewport={"width": 1440, "height": 900},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
        )
        page = context.new_page()
        try:
            downloader.login(page, debug_dir)
            downloader.open_module(page)
            for section in downloader.SECTIONS:
                result = downloader.process_one(page, debug_dir, section, target)
                results.append(result)
                logging.info("%s | %s | %s", result["Section"], result["Month"], result["Status"])
        finally:
            browser.close()

    downloader.write_log(results)
    failed = [r for r in results if r["Status"] == "FAILED"]
    if failed:
        names = ", ".join(str(r["Section"]) for r in failed)
        raise RuntimeError(f"Current-month download failed for: {names}")
    return results


def current_month_files() -> list[tuple[str, Path]]:
    month = datetime.now().strftime("%Y-%m")
    return [
        ("Return from Ops", DOWNLOAD_DIR / "1_Return_from_Ops" / f"Return_from_Ops_{month}.csv"),
        ("Handover to Ops", DOWNLOAD_DIR / "2_Handover_to_Ops" / f"Handover_to_Ops_{month}.csv"),
        ("Service Request", DOWNLOAD_DIR / "3_Service_Request" / f"Service_Request_{month}.csv"),
        ("Inter Station Transfer", DOWNLOAD_DIR / "4_Vehicle_Movement" / "Inter_Station_Transfer" / f"Inter_Station_Transfer_{month}.csv"),
        ("Vehicle Requisition", DOWNLOAD_DIR / "4_Vehicle_Movement" / "Vehicle_Requisition" / f"Vehicle_Requisition_{month}.csv"),
    ]


def resolve_file(path: Path) -> Path:
    if path.exists():
        return path
    # The Return export has previously used a vendor-supplied filename.
    month = datetime.now().strftime("%Y-%m")
    if path.name.startswith("Return_from_Ops_"):
        alternate = path.with_name(f"return_tickets_{month}.csv")
        if alternate.exists():
            return alternate
    raise FileNotFoundError(f"Expected current-month export is missing: {path}")


def read_sources() -> tuple[list[dict[str, str]], list[str]]:
    rows: list[dict[str, str]] = []
    source_columns: list[str] = []
    for source, expected in current_month_files():
        path = resolve_file(expected)
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            for column in reader.fieldnames or []:
                if column not in source_columns:
                    source_columns.append(column)
            for record in reader:
                clean = {str(k).strip(): (v or "").strip() for k, v in record.items() if k is not None}
                clean["Ticket Source"] = source
                clean["Source File"] = path.name
                rows.append(clean)
    return rows, source_columns


def first_value(row: dict[str, str], *names: str) -> str:
    return next((row.get(name, "") for name in names if row.get(name, "")), "")


def normalize(row: dict[str, str]) -> dict[str, str]:
    status = first_value(row, "Status", "Ticket Status")
    status_key = status.strip().casefold()
    accepted_on = first_value(row, "Accepted On")
    accepted_by = first_value(row, "Accepted By")
    closed_on = first_value(row, "Closed On")

    # Service Request exports call the latest acceptance/closure fields
    # "Last Action" rather than "Accepted" or "Closed".
    if status_key in COMPLETED_STATUSES:
        if not accepted_on or accepted_on == "-":
            accepted_on = first_value(row, "Last Action On")
        if not accepted_by or accepted_by == "-":
            accepted_by = first_value(row, "Last Action By")

    return {
        "Ticket Source": row["Ticket Source"],
        "Ticket ID": first_value(row, "Ticket ID", "Ticket Number"),
        "Accepted On": accepted_on,
        "Raised On": first_value(row, "Raised On", "Created On"),
        "Closed On": closed_on,
        "Accepted By": accepted_by,
        "Vehicle No": first_value(row, "Vehicle No", "Vehicle"),
        "Site": first_value(row, "Site Code", "Required Site Code", "From Site"),
        "Assigned To": first_value(row, "Assigned To", "From OM"),
        "Status": status,
        "Age (Days)": first_value(row, "Age (Days)", "Ageing (Days)"),
        "Reason / Remarks": first_value(row, "Reason", "Remarks"),
        "Raised By": first_value(row, "Raised By"),
        "Source File": row["Source File"],
    }


def ticket_key(row: dict[str, str]) -> tuple[str, str]:
    """Return a stable key used to compare daily completed-ticket snapshots."""
    return (row.get("Ticket Source", "").strip().casefold(), row.get("Ticket ID", "").strip().casefold())


def previous_report() -> Path | None:
    """Find the newest dated report before today, preferring the actual prior day."""
    candidates: list[tuple[date, Path]] = []
    pattern = re.compile(r"Current_Month_Open_Tickets_(\d{2}-\d{2}-\d{2})\.xlsx$", re.IGNORECASE)
    for path in REPORT_DIR.glob("Current_Month_Open_Tickets_*.xlsx"):
        match = pattern.match(path.name)
        if not match:
            continue
        try:
            report_date = datetime.strptime(match.group(1), "%d-%m-%y").date()
        except ValueError:
            continue
        if report_date < datetime.now().date():
            candidates.append((report_date, path))
    return max(candidates, default=(None, None), key=lambda item: item[0])[1]


def previous_completed_ticket_keys() -> tuple[set[tuple[str, str]], Path | None]:
    path = previous_report()
    if path is None:
        return set(), None

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "Completed Tickets" not in wb.sheetnames:
            wb.close()
            return set(), path
        ws = wb["Completed Tickets"]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        source_index = headers.index("Ticket Source") if "Ticket Source" in headers else -1
        ticket_index = headers.index("Ticket ID") if "Ticket ID" in headers else -1
        keys: set[tuple[str, str]] = set()
        if source_index >= 0 and ticket_index >= 0:
            for values in rows:
                source = str(values[source_index] or "").strip().casefold()
                ticket_id = str(values[ticket_index] or "").strip().casefold()
                if ticket_id:
                    keys.add((source, ticket_id))
        wb.close()
        return keys, path
    except Exception as exc:
        logging.warning("Could not read previous report %s: %s", path, exc)
        return set(), path


def style_header(ws: Any, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_table(ws: Any, start_row: int, end_row: int, end_col: int, name: str) -> None:
    if end_row <= start_row:
        return
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def autosize(ws: Any, maximum: int = 38) -> None:
    for column in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(r, column).value or "")) for r in range(1, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 11), maximum)


def build_report(rows: list[dict[str, str]], source_columns: list[str]) -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    normalized = [normalize(row) for row in rows]
    open_rows = [row for row in normalized if row["Status"].strip().casefold() == "open"]
    completed_rows = [row for row in normalized if row["Status"].strip().casefold() in COMPLETED_STATUSES]
    previous_keys, previous_path = previous_completed_ticket_keys()
    newly_completed_rows = [] if previous_path is None else [
        row for row in completed_rows
        if row["Ticket ID"].strip() and ticket_key(row) not in previous_keys
    ]
    source_names = [source for source, _ in current_month_files()]

    wb = Workbook()
    summary = wb.active
    summary.title = "Open Ticket Summary"
    summary.sheet_view.showGridLines = False
    summary["A1"] = "CURRENT MONTH OPEN TICKETS"
    summary["A1"].font = Font(size=18, bold=True, color=WHITE)
    summary["A1"].fill = PatternFill("solid", fgColor=BLUE)
    summary.merge_cells("A1:E1")
    summary["A3"] = "Open tickets at present"
    summary["B3"] = len(open_rows)
    summary["A3"].font = Font(size=14, bold=True)
    summary["B3"].font = Font(size=20, bold=True, color=RED)
    summary["A4"] = "Total completed tickets this month"
    summary["B4"] = len(completed_rows)
    summary["A4"].font = Font(size=14, bold=True)
    summary["B4"].font = Font(size=20, bold=True, color=GREEN)
    summary["A5"] = "Completed tickets in previous report"
    summary["B5"] = len(previous_keys)
    summary["A6"] = "Newly completed since previous report"
    summary["B6"] = len(newly_completed_rows)
    summary["B6"].font = Font(size=20, bold=True, color=GREEN)
    summary["A7"] = "Compared with"
    summary["B7"] = previous_path.name if previous_path else "No earlier daily report found"
    summary["A8"] = "Report month"
    summary["B8"] = datetime.now().strftime("%B %Y")
    summary["A9"] = "Generated on"
    summary["B9"] = datetime.now().strftime("%d %b %Y %I:%M %p")
    summary.append([])
    summary.append(["Ticket Source", "Open", "Completed", "Newly Completed", "All Current-Month Rows", "Action"])
    style_header(summary, 11, 6)
    for source in source_names:
        total = sum(row["Ticket Source"] == source for row in normalized)
        opened = sum(row["Ticket Source"] == source for row in open_rows)
        completed = sum(row["Ticket Source"] == source for row in completed_rows)
        newly_completed = sum(row["Ticket Source"] == source for row in newly_completed_rows)
        summary.append([source, opened, completed, newly_completed, total, "Close pending tickets" if opened else "No open tickets"])
    total_row = summary.max_row + 1
    summary.append(["TOTAL", len(open_rows), len(completed_rows), len(newly_completed_rows), len(normalized), ""])
    for cell in summary[total_row]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    summary.conditional_formatting.add(
        f"B12:B{total_row - 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_RED))
    )
    summary.conditional_formatting.add(
        f"C12:D{total_row - 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_BLUE))
    )
    summary.freeze_panes = "A12"
    autosize(summary)

    completed_ws = wb.create_sheet("Completed Tickets")
    normalized_headers = list(normalized[0].keys()) if normalized else list(normalize({"Ticket Source": "", "Source File": ""}).keys())
    completed_headers = normalized_headers
    completed_ws.append(completed_headers)
    for row in completed_rows:
        completed_ws.append([row.get(header, "") for header in completed_headers])
    style_header(completed_ws, 1, len(completed_headers))
    completed_ws.freeze_panes = "A2"
    completed_ws.auto_filter.ref = completed_ws.dimensions
    add_table(completed_ws, 1, completed_ws.max_row, len(completed_headers), "CompletedTicketsTable")
    autosize(completed_ws)

    completed_today_ws = wb.create_sheet("Completed Today")
    completed_today_ws.append(completed_headers)
    for row in newly_completed_rows:
        completed_today_ws.append([row.get(header, "") for header in completed_headers])
    style_header(completed_today_ws, 1, len(completed_headers))
    completed_today_ws.freeze_panes = "A2"
    completed_today_ws.auto_filter.ref = completed_today_ws.dimensions
    add_table(
        completed_today_ws,
        1,
        completed_today_ws.max_row,
        len(completed_headers),
        "CompletedTodayTable",
    )
    autosize(completed_today_ws)

    open_ws = wb.create_sheet("Open Tickets")
    open_headers = normalized_headers
    open_ws.append(open_headers)
    for row in open_rows:
        open_ws.append([row.get(header, "") for header in open_headers])
    style_header(open_ws, 1, len(open_headers))
    open_ws.freeze_panes = "A2"
    open_ws.auto_filter.ref = open_ws.dimensions
    add_table(open_ws, 1, open_ws.max_row, len(open_headers), "OpenTicketsTable")
    autosize(open_ws)

    all_ws = wb.create_sheet("All Current Month")
    all_headers = ["Ticket Source", *source_columns, "Source File"]
    all_ws.append(all_headers)
    for row in rows:
        all_ws.append([row.get(header, "") for header in all_headers])
    style_header(all_ws, 1, len(all_headers))
    all_ws.freeze_panes = "A2"
    all_ws.auto_filter.ref = all_ws.dimensions
    add_table(all_ws, 1, all_ws.max_row, len(all_headers), "AllTicketsTable")
    autosize(all_ws)

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(REPORT_PATH)
    logging.info(
        "Report saved: %s (%d open, %d completed, %d newly completed)",
        REPORT_PATH,
        len(open_rows),
        len(completed_rows),
        len(newly_completed_rows),
    )
    return REPORT_PATH


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report-only", action="store_true", help="Use existing current-month CSVs without downloading.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    log_path = setup_logging()
    try:
        if not args.report_only:
            download_current_month()
        rows, columns = read_sources()
        report = build_report(rows, columns)
        print(f"Created: {report}")
        print(f"Log: {log_path}")
        return 0
    except Exception:
        logging.exception("Daily open-ticket job failed")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
